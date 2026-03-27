"""Excel (.xlsx) to Markdown converter — merged cells + auto data range."""

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .converter_base import ConverterBase


class ExcelConverter(ConverterBase):
    """Convert .xlsx to Markdown with merged cell handling."""

    def convert(self, filepath: str) -> str:
        wb = load_workbook(filepath, data_only=True)
        sections = []

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            section = self._convert_sheet(ws, sheet_name)
            if section:
                sections.append(section)

        result = "\n\n---\n\n".join(sections)
        return self.normalize_markdown(self.cleanup_text(result))

    def _convert_sheet(self, ws, sheet_name: str) -> str:
        """Convert a single worksheet to markdown."""
        # Build merged cells lookup: (row, col) -> master cell value
        merged_lookup = {}
        merged_masters = set()
        for merged_range in ws.merged_cells.ranges:
            master_cell = ws.cell(merged_range.min_row, merged_range.min_col)
            master_val = self._format_cell(master_cell)
            merged_masters.add((merged_range.min_row, merged_range.min_col))
            for row in range(merged_range.min_row, merged_range.max_row + 1):
                for col in range(merged_range.min_col, merged_range.max_col + 1):
                    if (row, col) != (merged_range.min_row, merged_range.min_col):
                        merged_lookup[(row, col)] = ""  # Empty for non-master
                    else:
                        merged_lookup[(row, col)] = master_val

        # Detect actual data range
        min_row = ws.min_row or 1
        max_row = ws.max_row or 1
        min_col = ws.min_column or 1
        max_col = ws.max_column or 1

        # Skip trailing empty rows/cols
        while max_row > min_row:
            if any(ws.cell(max_row, c).value for c in range(min_col, max_col + 1)):
                break
            max_row -= 1

        while max_col > min_col:
            if any(ws.cell(r, max_col).value for r in range(min_row, max_row + 1)):
                break
            max_col -= 1

        if max_row < min_row or max_col < min_col:
            return ""

        # Build table data
        rows_data = []
        for row in range(min_row, max_row + 1):
            row_cells = []
            for col in range(min_col, max_col + 1):
                if (row, col) in merged_lookup:
                    row_cells.append(merged_lookup[(row, col)])
                else:
                    cell = ws.cell(row, col)
                    row_cells.append(self._format_cell(cell))
            rows_data.append(row_cells)

        if not rows_data:
            return ""

        # Build markdown table
        num_cols = max_col - min_col + 1
        lines = [f"## 📊 {sheet_name}", ""]

        # Header row
        header = "| " + " | ".join(rows_data[0]) + " |"
        separator = "| " + " | ".join(["---"] * num_cols) + " |"
        lines.append(header)
        lines.append(separator)

        # Data rows
        for row in rows_data[1:]:
            lines.append("| " + " | ".join(row) + " |")

        return "\n".join(lines)

    @staticmethod
    def _format_cell(cell) -> str:
        """Format cell value for markdown."""
        val = cell.value
        if val is None:
            return ""

        # Format numbers
        if isinstance(val, float):
            if val == int(val):
                return str(int(val))
            return f"{val:,.2f}"
        elif isinstance(val, int):
            return f"{val:,}"

        # Format dates
        if hasattr(val, "strftime"):
            try:
                return val.strftime("%Y-%m-%d")
            except Exception:
                pass

        text = str(val).replace("|", "\\|").replace("\n", " ").strip()
        return text
