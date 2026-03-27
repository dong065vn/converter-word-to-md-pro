"""PowerPoint (.pptx) to Markdown converter — recursive shape traversal."""

import os
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from .converter_base import ConverterBase


class PptxConverter(ConverterBase):
    """Convert .pptx to Markdown with recursive shape handling."""

    def convert(self, filepath: str) -> str:
        prs = Presentation(filepath)
        slides_md = []

        for slide_num, slide in enumerate(prs.slides, 1):
            slide_lines = []

            # Slide title
            title = ""
            if slide.shapes.title:
                title = slide.shapes.title.text.strip()
            if title:
                slide_lines.append(f"## Slide {slide_num}: {title}")
            else:
                slide_lines.append(f"## Slide {slide_num}")
            slide_lines.append("")

            # Process shapes recursively
            for shape in slide.shapes:
                if shape == slide.shapes.title:
                    continue  # Already handled
                shape_md = self._process_shape(shape, filepath)
                if shape_md:
                    slide_lines.append(shape_md)

            # Speaker notes
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes_text = slide.notes_slide.notes_text_frame.text.strip()
                if notes_text:
                    slide_lines.append("")
                    slide_lines.append(f"> **Notes:** {notes_text}")

            slides_md.append("\n".join(slide_lines))

        result = "\n\n---\n\n".join(slides_md)
        return self.normalize_markdown(self.cleanup_text(result))

    def _process_shape(self, shape, filepath: str) -> str:
        """Process a shape recursively (handles GroupShape)."""
        parts = []

        # Group shape — recurse
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            for child_shape in shape.shapes:
                child_md = self._process_shape(child_shape, filepath)
                if child_md:
                    parts.append(child_md)

        # Table
        elif shape.has_table:
            table_md = self._process_table(shape.table)
            if table_md:
                parts.append(table_md)

        # Text frame
        elif shape.has_text_frame:
            text_md = self._process_text_frame(shape.text_frame)
            if text_md:
                parts.append(text_md)

        # Image
        if hasattr(shape, "image") and shape.image:
            try:
                img_md = self._process_image(shape, filepath)
                if img_md:
                    parts.append(img_md)
            except Exception:
                pass

        return "\n".join(parts)

    def _process_text_frame(self, text_frame) -> str:
        """Process text frame with formatting."""
        lines = []
        for para in text_frame.paragraphs:
            text = ""
            for run in para.runs:
                run_text = run.text or ""
                if not run_text.strip():
                    text += run_text
                    continue
                if run.font.bold and run.font.italic:
                    run_text = f"***{run_text}***"
                elif run.font.bold:
                    run_text = f"**{run_text}**"
                elif run.font.italic:
                    run_text = f"*{run_text}*"
                text += run_text

            text = text.strip()
            if not text:
                continue

            # Detect bullet level
            level = para.level or 0
            if level > 0:
                indent = "  " * (level - 1)
                text = f"{indent}- {text}"
            elif para.level == 0 and len(para.runs) > 0:
                # Check if it looks like a sub-heading (bold, short)
                if all(r.font.bold for r in para.runs if r.text.strip()) and len(text) < 80:
                    text = f"### {text.replace('**', '').replace('***', '')}"

            lines.append(text)

        return "\n".join(lines) if lines else ""

    def _process_table(self, table) -> str:
        """Process PowerPoint table to markdown."""
        rows_data = []
        for row in table.rows:
            cells = []
            for cell in row.cells:
                cell_text = cell.text.replace("|", "\\|").replace("\n", " ").strip()
                cells.append(cell_text)
            rows_data.append(cells)

        if not rows_data:
            return ""

        max_cols = max(len(r) for r in rows_data)
        for row in rows_data:
            while len(row) < max_cols:
                row.append("")

        lines = []
        lines.append("| " + " | ".join(rows_data[0]) + " |")
        lines.append("| " + " | ".join(["---"] * max_cols) + " |")
        for row in rows_data[1:]:
            lines.append("| " + " | ".join(row) + " |")

        return "\n".join(lines)

    def _process_image(self, shape, filepath: str) -> str:
        """Extract and save image from shape."""
        try:
            image = shape.image
            img_bytes = image.blob
            ext = image.content_type.split("/")[-1]
            if ext == "jpeg":
                ext = "jpg"

            img_name = f"slide_img_{id(shape)}.{ext}"
            img_path = os.path.join(self.image_output_dir, img_name)
            os.makedirs(os.path.dirname(img_path), exist_ok=True)

            with open(img_path, "wb") as f:
                f.write(img_bytes)

            return f"![image]({img_path})"
        except Exception:
            return ""
